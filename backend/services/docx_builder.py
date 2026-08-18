import re
import uuid
from pathlib import Path

from config import settings
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_cell_background(cell, color: str):
    """Set cell background color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def add_heading(doc: Document, text: str, level: int):
    """Add heading with consistent styling."""
    heading = doc.add_heading(text, level=level)
    run = heading.runs[0] if heading.runs else heading.add_run(text)
    run.font.name = "Arial"
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)
    return heading


def add_paragraph(doc: Document, text: str):
    """Add standard paragraph."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    return p


def add_cover_page(doc: Document, device_data: dict):
    """Build the document cover page."""
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Software Architectural Design — Draft for Review")
    run.font.name = "Arial"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    doc.add_paragraph()

    # Device info
    info_lines = [
        f"Device Name: {device_data['name']}",
        f"Software Safety Class: {device_data['safety_class']}",
        f"Driver Software Version: {device_data.get('driver_version', '01')}",
        f"GUI Software Version: {device_data.get('gui_version', '1.0.0')}",
    ]
    for line in info_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.font.name = "Arial"
        run.font.size = Pt(12)

    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice_run = notice.add_run(
        "AI-assisted draft. Technical, quality, and regulatory review is required before release."
    )
    notice_run.font.name = "Arial"
    notice_run.font.size = Pt(10)
    notice_run.font.italic = True

    doc.add_paragraph()

    # Signature table
    table = doc.add_table(rows=2, cols=6)
    table.style = "Table Grid"
    headers = ["Signature", "Date", "Approved by", "Confirmed by", "Prepared by", "Rev."]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "D5E8F0")

    doc.add_page_break()


def add_alarms_table(doc: Document, alarms: list):
    """Alarms and warnings table."""
    if not alarms:
        add_paragraph(doc, "No alarms defined for this device.")
        return

    headers = ["No.", "Priority", "Condition", "Text Shown",
               "Indicator", "Sound", "Required Action"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    # header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "D5E8F0")

    # data rows
    for j, alarm in enumerate(alarms):
        row = table.add_row()
        values = [
            str(j + 1),
            alarm.get("priority", ""),
            alarm.get("condition", ""),
            alarm.get("text_shown", ""),
            alarm.get("indicator_light", ""),
            "Yes" if alarm.get("indicator_sound") else "No",
            alarm.get("required_action", ""),
        ]
        for i, value in enumerate(values):
            row.cells[i].text = value or ""


def build_document(output_data: dict) -> str:
    """Build the final Word document."""
    device_data = output_data["device_data"]
    sections = output_data["sections"]

    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

# Cover page
    add_cover_page(doc, device_data)

    # Preserve every generated section in document order.
    for section_name, content in sections.items():
        match = re.match(r"^\s*(\d+(?:\.\d+)*)", section_name)
        level = min(match.group(1).count(".") + 1, 3) if match else 1
        add_heading(doc, section_name, level=level)
        add_paragraph(doc, content)

        if "gui software" in section_name.lower() and device_data.get("alarms"):
            add_heading(doc, "Errors and Warnings Display", level=min(level + 1, 3))
            add_alarms_table(doc, device_data["alarms"])

    # Save file
    output_path = Path(settings.generated_docs_path)
    output_path.mkdir(parents=True, exist_ok=True)

    file_id = str(uuid.uuid4())[:8]
    device_name = device_data["name"].replace(" ", "_")
    filename = f"{device_name}_{file_id}.docx"
    filepath = output_path / filename

    doc.save(str(filepath))
    return str(filepath)
