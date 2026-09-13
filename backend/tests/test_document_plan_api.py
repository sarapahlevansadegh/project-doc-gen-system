"""Integration test for GET /rag/reference/{reference_id}/plan - verifies
the actual HTTP endpoint wiring for services/document_diff_planner.py, not
just the service function directly (test_document_diff_planner.py already
covers build_document_plan()/build_section_plan() at that level in depth).

Same hermetic spirit as test_journey.py: no real LLM or embedding model,
both are monkeypatched out. Reference section embedding and device chunk
embedding are set to the SAME fixed non-zero vector so pgvector's cosine
distance resolves to a real, deterministic similarity (1.0) above
DEFAULT_SIMILARITY_THRESHOLD - this is what exercises the LLM-call branch
of build_section_plan, not just the "no match" early-return branch.
"""
from __future__ import annotations

import asyncio
import io
import json
import uuid

import pytest
from conftest import (
    cleanup_seeded_user,
    dispose_engine_sync,
    login_headers_sync,
    seed_admin_user,
)
from docx import Document as DocxDocument

EMBED_DIM = 768
_FIXED_VECTOR = [1.0] + [0.0] * (EMBED_DIM - 1)


def _fake_llm_generate(self, prompt: str, max_tokens: int = 800, temperature: float = 0.1) -> str:
    # document_diff_planner's section prompt always includes this heading.
    if "REFERENCE SECTION" in prompt:
        return json.dumps(
            {
                "changed": False,
                "reason": "test stub - no device-specific replacement needed",
                "new_paragraphs": None,
                "new_table_markdown": None,
            }
        )
    # ontology/extractor.py's relationship prompt always embeds "ENTITIES:";
    # keep both ontology calls (entities + relationships) trivially empty -
    # this test is about the /plan endpoint wiring, not ontology content.
    return "[]"


async def _fake_embed_text(text: str) -> list[float]:
    return list(_FIXED_VECTOR)


class _FakeDeviceEmbeddingService:
    def embed_documents(self, texts: list[str], batch_size: int = 32) -> list[list[float]]:
        return [list(_FIXED_VECTOR) for _ in texts]


def _make_reference_docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("1. Architecture of Software", level=1)
    doc.add_paragraph("The device runs a layered software architecture.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_device_docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("1. Architecture of Software", level=1)
    doc.add_paragraph("The High Power Laser VL8 device runs a layered software architecture.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _db_ok(AsyncSessionLocal):
    from sqlalchemy import text

    async with AsyncSessionLocal() as db:
        await db.execute(text("SELECT 1"))


def test_get_document_plan_endpoint_returns_section_plans_and_ontology_warnings():
    try:
        from agent.llm import LLMClient
        from app import app
        from fastapi.testclient import TestClient
        from rag import chunk_service
        from rag import retriever as rag_retriever
    except ModuleNotFoundError:
        pytest.skip("async app stack / TestClient unavailable")

    from core.database import AsyncSessionLocal, get_engine

    loop = asyncio.new_event_loop()
    admin_creds = None
    try:
        try:
            loop.run_until_complete(_db_ok(AsyncSessionLocal))
        except Exception as exc:
            if "SELECT 1" in str(exc) or "connect" in str(exc).lower():
                pytest.skip("Postgres unavailable")
            raise
        admin_creds = seed_admin_user(loop)
    finally:
        try:
            loop.run_until_complete(get_engine().dispose())
        except Exception:
            pass
        loop.close()

    reference_id: str | None = None
    device_id: str | None = None
    try:
        with TestClient(app) as client:
            client.headers.update(login_headers_sync(client, *admin_creds))

            real_generate = LLMClient.generate
            real_embed_text = rag_retriever.embed_text
            real_device_embed_service = chunk_service.get_device_embedding_service
            LLMClient.generate = _fake_llm_generate
            rag_retriever.embed_text = _fake_embed_text
            chunk_service.get_device_embedding_service = lambda: _FakeDeviceEmbeddingService()
            try:
                # 1. Upload the reference document
                ref_resp = client.post(
                    "/rag/reference",
                    files={
                        "file": (
                            "reference.docx",
                            _make_reference_docx(),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    },
                )
                assert ref_resp.status_code == 200, ref_resp.text
                reference_id = ref_resp.json()["id"]

                # 2. Upload the device document (chunks + ontology now run
                # automatically as part of this call - see
                # services/device_document_service.py's _attach_document)
                dev_resp = client.post(
                    "/devices/documents",
                    files={
                        "file": (
                            "device.docx",
                            _make_device_docx(),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    },
                )
                assert dev_resp.status_code == 201, dev_resp.text
                device = dev_resp.json()
                device_id = device["id"]
                device_document_id = device["documents"][0]["id"]

                # 3. Call the actual HTTP endpoint under test
                plan_resp = client.get(
                    f"/rag/reference/{reference_id}/plan",
                    params={"device_document_id": device_document_id},
                )
                assert plan_resp.status_code == 200, plan_resp.text
                body = plan_resp.json()

                assert "section_plans" in body and "ontology_warnings" in body
                assert len(body["section_plans"]) >= 1
                section_plan = body["section_plans"][0]
                assert section_plan["section_name"]
                # similarity was forced to 1.0 (identical fixed vectors),
                # well above DEFAULT_SIMILARITY_THRESHOLD (0.5) - confirms
                # the LLM branch actually ran, not the "no match" fallback
                assert section_plan["best_similarity"] == pytest.approx(1.0, abs=1e-6)
                assert section_plan["reason"] == "test stub - no device-specific replacement needed"
                assert isinstance(body["ontology_warnings"], list)
            finally:
                LLMClient.generate = real_generate
                rag_retriever.embed_text = real_embed_text
                chunk_service.get_device_embedding_service = real_device_embed_service
    finally:
        if admin_creds:
            cleanup_seeded_user(admin_creds[0])

        async def _cleanup():
            from models.device import Device
            from models.ontology import OntologyEntity
            from models.template import DocumentTemplate, ReferenceDocument

            async with AsyncSessionLocal() as db:
                if reference_id:
                    await db.execute(
                        DocumentTemplate.__table__.delete().where(
                            DocumentTemplate.source_doc_id == uuid.UUID(reference_id)
                        )
                    )
                    await db.execute(
                        ReferenceDocument.__table__.delete().where(
                            ReferenceDocument.id == uuid.UUID(reference_id)
                        )
                    )
                if device_id:
                    await db.execute(
                        OntologyEntity.__table__.delete().where(
                            OntologyEntity.device_id == uuid.UUID(device_id)
                        )
                    )
                    dev = await db.get(Device, uuid.UUID(device_id))
                    if dev is not None:
                        await db.delete(dev)
                await db.commit()

        try:
            asyncio.run(_cleanup())
        except Exception:
            pass

    dispose_engine_sync()
