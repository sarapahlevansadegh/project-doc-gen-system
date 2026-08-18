"""Reference document management tests (Milestone 6.2)."""
from __future__ import annotations

import asyncio
import io

import pytest
from docx import Document as DocxDocument

EMBED_DIM = 384


async def _fake_embed(text: str) -> list[float]:
    return [0.0] * EMBED_DIM


def _make_reference_docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("1. Architecture of Software", level=1)
    doc.add_paragraph("The device runs a layered software architecture.")
    doc.add_heading("1.1 LCD Module", level=2)
    doc.add_paragraph("The LCD module renders status and warnings.")
    doc.add_heading("2. Serial Communication Architecture", level=1)
    doc.add_paragraph("Serial link uses RS-232 at 115200 baud.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def _client():
    try:
        from app import app
        from core.database import AsyncSessionLocal, get_engine
        from fastapi.testclient import TestClient
        from rag import retriever as rag_retriever
    except ModuleNotFoundError:
        pytest.skip("async app stack / TestClient unavailable")

    loop = asyncio.new_event_loop()
    try:
        from sqlalchemy import text

        async def _ok():
            async with AsyncSessionLocal() as db:
                await db.execute(text("SELECT 1"))

        try:
            loop.run_until_complete(_ok())
        except Exception as exc:
            if "SELECT 1" in str(exc) or "connect" in str(exc).lower():
                pytest.skip("Postgres unavailable")
            raise
    finally:
        try:
            loop.run_until_complete(get_engine().dispose())
        except Exception:
            pass
        loop.close()

    with TestClient(app) as client:
        with _Monkey(rag_retriever, "embed_text", _fake_embed):
            yield client


class _Monkey:
    def __init__(self, obj, attr, fake):
        self._obj = obj
        self._attr = attr
        self._fake = fake
        self._real = None

    def __enter__(self):
        self._real = getattr(self._obj, self._attr)
        setattr(self._obj, self._attr, self._fake)
        return self

    def __exit__(self, *exc):
        setattr(self._obj, self._attr, self._real)


def _upload(client):
    resp = client.post(
        "/rag/reference",
        files={"file": ("ref.docx", _make_reference_docx(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_list_includes_created_at(_client):
    ref = _upload(_client)
    listing = _client.get("/rag/reference").json()
    assert any(r["id"] == ref["id"] for r in listing)
    match = next(r for r in listing if r["id"] == ref["id"])
    assert match["section_count"] >= 1
    assert "created_at" in match
    assert match["created_at"] is not None
    assert "is_active" in match


def test_activate_enforces_single_active(_client):
    a = _upload(_client)
    b = _upload(_client)

    # activate b
    act = _client.post(f"/rag/reference/{b['id']}/activate")
    assert act.status_code == 200, act.text
    assert act.json()["is_active"] is True

    listing = {r["id"]: r for r in _client.get("/rag/reference").json()}
    assert listing[b["id"]]["is_active"] is True
    assert listing[a["id"]]["is_active"] is False

    # active endpoint returns b
    active = _client.get("/rag/reference/active").json()
    assert active["id"] == b["id"]


def test_activate_missing_returns_404(_client):
    import uuid

    resp = _client.post(f"/rag/reference/{uuid.uuid4()}/activate")
    assert resp.status_code == 404


def test_delete_removes_reference_and_sections(_client):
    ref = _upload(_client)
    ref_id = ref["id"]

    # sections exist
    sections = _client.get(f"/rag/reference/{ref_id}/sections").json()
    assert len(sections) >= 1

    del_resp = _client.delete(f"/rag/reference/{ref_id}")
    assert del_resp.status_code == 204

    # sections removed (query returns empty list for the deleted reference)
    assert _client.get(f"/rag/reference/{ref_id}/sections").json() == []
    assert all(r["id"] != ref_id for r in _client.get("/rag/reference").json())


def test_delete_missing_returns_404(_client):
    import uuid

    assert _client.delete(f"/rag/reference/{uuid.uuid4()}").status_code == 404
