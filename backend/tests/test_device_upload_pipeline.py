"""Integration test for the Phase 3.5.9->upload wiring: chunk generation
(Phase 3) and ontology extraction (Phase 3.5) must now run automatically as
part of POST /devices/documents, not just be available as separate manual
/chunks and /ontology endpoints.

Same hermetic spirit as test_journey.py: no real LLM or embedding model,
both are monkeypatched out. Uses TestClient (real HTTP layer) so this
actually exercises services/device_document_service.py's _attach_document
wiring, not just the underlying service functions in isolation (those are
already covered by test_chunk_service.py and
test_ontology_extraction_service.py).
"""
from __future__ import annotations

import io
import json
import uuid

import pytest
from conftest import (
    cleanup_seeded_user,
    dispose_engine_sync,
    login_headers_sync,
    seed_admin_user,
)
from docx import Document as DocxDocument

EMBED_DIM = 768


def _fake_llm_generate(self, prompt: str, max_tokens: int = 800, temperature: float = 0.1) -> str:
    # ontology/extractor.py's relationship prompt always embeds an
    # "ENTITIES:" section - same discriminator test_ontology_extraction_
    # service.py's _FakeLLM uses to tell the two prompt shapes apart.
    if "ENTITIES:" in prompt:
        return "[]"
    return json.dumps([{"type": "Device", "value": "VL8"}])


class _FakeEmbeddingService:
    def embed_documents(self, texts: list[str], batch_size: int = 32) -> list[list[float]]:
        return [[0.0] * EMBED_DIM for _ in texts]


def _make_device_docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("1. Architecture of Software", level=1)
    doc.add_paragraph("The High Power Laser VL8 device runs a layered software architecture.")
    doc.add_heading("2. Driver Software Version", level=1)
    doc.add_paragraph("Driver Software Version: 01")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _db_ok(AsyncSessionLocal):
    from sqlalchemy import text

    async with AsyncSessionLocal() as db:
        await db.execute(text("SELECT 1"))


def test_upload_automatically_generates_chunks_and_ontology():
    try:
        from agent.llm import LLMClient
        from app import app
        from fastapi.testclient import TestClient
        from rag import chunk_service
    except ModuleNotFoundError:
        pytest.skip("async app stack / TestClient unavailable")

    import asyncio

    from core.database import AsyncSessionLocal, get_engine

    loop = asyncio.new_event_loop()
    admin_creds = None
    try:
        try:
            loop.run_until_complete(_db_ok(AsyncSessionLocal))
        except Exception as exc:
            if "SELECT 1" in str(exc) or "connect" in str(exc).lower():
                pytest.skip("Postgres unavailable")
            raise
        admin_creds = seed_admin_user(loop)
    finally:
        try:
            loop.run_until_complete(get_engine().dispose())
        except Exception:
            pass
        loop.close()

    device_id: str | None = None
    try:
        with TestClient(app) as client:
            client.headers.update(login_headers_sync(client, *admin_creds))

            real_generate = LLMClient.generate
            real_embed_service = chunk_service.get_device_embedding_service
            LLMClient.generate = _fake_llm_generate
            chunk_service.get_device_embedding_service = lambda: _FakeEmbeddingService()
            try:
                docx_bytes = _make_device_docx()
                resp = client.post(
                    "/devices/documents",
                    files={
                        "file": (
                            "vl8-test.docx",
                            docx_bytes,
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    },
                )
                assert resp.status_code == 201, resp.text
                device = resp.json()
                device_id = device["id"]
                assert len(device["documents"]) == 1
                document_id = uuid.UUID(device["documents"][0]["id"])

                # Chunks: created automatically, no manual POST .../chunks call
                chunks_resp = client.get(f"/devices/documents/{document_id}/chunks")
                assert chunks_resp.status_code == 200, chunks_resp.text
                chunks = chunks_resp.json()
                assert len(chunks) >= 1
                assert all(c["has_embedding"] for c in chunks)

                # Ontology: created automatically, no manual POST .../ontology call
                async def _fetch_entities():
                    from models.ontology import OntologyEntity
                    from sqlalchemy import select

                    async with AsyncSessionLocal() as db:
                        result = await db.execute(
                            select(OntologyEntity).where(
                                OntologyEntity.device_id == uuid.UUID(device_id)
                            )
                        )
                        return result.scalars().all()

                entities = asyncio.run(_fetch_entities())
                assert any(
                    e.entity_type == "Device" and e.value == "VL8" for e in entities
                )
            finally:
                LLMClient.generate = real_generate
                chunk_service.get_device_embedding_service = real_embed_service
    finally:
        if admin_creds:
            cleanup_seeded_user(admin_creds[0])
        if device_id:
            async def _cleanup():
                from models.device import Device
                from models.ontology import OntologyEntity

                async with AsyncSessionLocal() as db:
                    await db.execute(
                        OntologyEntity.__table__.delete().where(
                            OntologyEntity.device_id == uuid.UUID(device_id)
                        )
                    )
                    dev = await db.get(Device, uuid.UUID(device_id))
                    if dev is not None:
                        await db.delete(dev)
                    await db.commit()

            try:
                asyncio.run(_cleanup())
            except Exception:
                pass

    dispose_engine_sync()
