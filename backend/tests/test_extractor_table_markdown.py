"""Unit tests for Markdown table serialization in rag/extractor.py.

Pure function tests (no DB, no FastAPI app) covering the flat-text -> real
Markdown table change.
"""
from __future__ import annotations

import io

from docx import Document as DocxDocument

from rag.extractor import _table_to_markdown, extract_docx


def test_table_to_markdown_basic():
    rows = [
        ["Parameter", "Value", "Unit"],
        ["Sample Rate", "250", "Hz"],
        ["Resolution", "12", "bit"],
    ]
    md = _table_to_markdown(rows)
    lines = md.splitlines()

    assert lines[0] == "| Parameter | Value | Unit |"
    assert lines[1] == "| --- | --- | --- |"
    assert lines[2] == "| Sample Rate | 250 | Hz |"
    assert lines[3] == "| Resolution | 12 | bit |"


def test_table_to_markdown_escapes_pipe_in_cell():
    rows = [["Field", "Allowed values"], ["Mode", "On|Off"]]
    md = _table_to_markdown(rows)
    assert "On\\|Off" in md


def test_table_to_markdown_pads_ragged_rows():
    rows = [["A", "B", "C"], ["only one"]]
    md = _table_to_markdown(rows)
    lines = md.splitlines()
    # ragged data row is padded to the header's column count
    assert lines[2] == "| only one |  |  |"


def test_table_to_markdown_empty_table():
    assert _table_to_markdown([]) == ""


def _make_docx_with_table() -> bytes:
    doc = DocxDocument()
    doc.add_heading("1. Specifications", level=1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Parameter"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Sample Rate"
    table.cell(1, 1).text = "250 Hz"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_docx_section_body_contains_markdown_table(tmp_path):
    path = tmp_path / "ref.docx"
    path.write_bytes(_make_docx_with_table())

    sections = extract_docx(path)
    spec_section = next(s for s in sections if s.title == "1. Specifications")

    assert "| Parameter | Value |" in spec_section.body
    assert "| --- | --- |" in spec_section.body
    assert spec_section.has_table is True
