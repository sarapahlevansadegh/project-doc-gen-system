"""Shape guard for the LLM's section plan, and the paragraph-count guard
when a plan is applied to a live document.

Everything here is database-free: the shape check is a pure function, and
the apply-time guard is exercised on a real python-docx document with
``build_section_plan`` and ``embed_text`` monkeypatched out.
"""
from __future__ import annotations

import uuid

import pytest
from docx import Document as DocxDocument
from services import document_generator
from services.document_diff_planner import SectionPlan, _plan_shape_problems


def _valid() -> dict:
    return {
        "changed": True,
        "reason": "updated",
        "new_paragraphs": ["one", "two"],
        "new_table_markdown": None,
    }


def test_valid_plans_have_no_problems():
    assert _plan_shape_problems(_valid()) == []
    # None means "not provided" for every optional field, and [] / "" are
    # legitimate empty values of the right type.
    assert _plan_shape_problems({"changed": False}) == []
    assert _plan_shape_problems({"changed": False, "new_paragraphs": [], "new_table_markdown": ""}) == []


@pytest.mark.parametrize("parsed", [[], "text", 0, 7, None, True])
def test_non_object_response_is_rejected(parsed):
    assert _plan_shape_problems(parsed) == ["response is not a JSON object"]


@pytest.mark.parametrize(
    "field, value, expected",
    [
        # falsy values of the wrong type must NOT slip through
        ("new_paragraphs", "", "new_paragraphs is not a list"),
        ("new_paragraphs", 0, "new_paragraphs is not a list"),
        ("new_paragraphs", False, "new_paragraphs is not a list"),
        ("new_paragraphs", {}, "new_paragraphs is not a list"),
        ("new_table_markdown", 0, "new_table_markdown is not a string"),
        ("new_table_markdown", [], "new_table_markdown is not a string"),
        ("new_table_markdown", {}, "new_table_markdown is not a string"),
        ("new_table_markdown", False, "new_table_markdown is not a string"),
        # truthy values of the wrong type
        ("new_paragraphs", "one paragraph", "new_paragraphs is not a list"),
        ("new_paragraphs", ["ok", 5], "new_paragraphs contains a non-string item"),
        ("new_paragraphs", ["ok", None], "new_paragraphs contains a non-string item"),
        ("new_table_markdown", ["| a |"], "new_table_markdown is not a string"),
        # "false" is truthy; 0/1 are not booleans
        ("changed", "false", "changed is not a boolean"),
        ("changed", 1, "changed is not a boolean"),
        ("changed", 0, "changed is not a boolean"),
        ("reason", 5, "reason is not a string"),
        ("reason", ["a"], "reason is not a string"),
    ],
)
def test_invalid_field_shapes_are_reported(field, value, expected):
    parsed = _valid()
    parsed[field] = value
    assert expected in _plan_shape_problems(parsed)


# --------------------------- apply-time guard ---------------------------


def _make_section(tmp_path, paragraph_texts):
    from rag.extractor import extract_docx_live

    doc = DocxDocument()
    doc.add_heading("Sensor Calibration", level=1)
    for text in paragraph_texts:
        doc.add_paragraph(text)
    path = tmp_path / "ref.docx"
    doc.save(str(path))
    _, sections = extract_docx_live(str(path))
    section = next(s for s in sections if "Sensor" in s.title)
    return section


@pytest.fixture
def stub_planner(monkeypatch):
    holder: dict = {}

    async def fake_embed(_text):
        return [0.0] * 768

    async def fake_plan(db, synthetic, device_document_id, llm, **kwargs):
        return holder["plan"]

    monkeypatch.setattr(document_generator, "embed_text", fake_embed)
    monkeypatch.setattr(document_generator, "build_section_plan", fake_plan)
    return holder


def _plan(new_paragraphs):
    return SectionPlan(
        section_id=uuid.uuid4(),
        section_name="Sensor Calibration",
        changed=True,
        reason="model reason",
        new_paragraphs=new_paragraphs,
    )


async def _apply(section):
    return await document_generator._plan_and_apply_section(
        None, section, uuid.uuid4(), None
    )


async def test_matching_paragraph_count_is_applied(tmp_path, stub_planner):
    section = _make_section(tmp_path, ["old one", "old two"])
    stub_planner["plan"] = _plan(["new one", "new two"])

    changed, reason = await _apply(section)

    texts = [b.text for b in section.blocks if hasattr(b, "text") and b.text.startswith(("new", "old"))]
    assert changed is True
    assert texts == ["new one", "new two"]
    assert "needs_review" not in reason


@pytest.mark.parametrize("new_paragraphs", [["only one"], ["a", "b", "c"]])
async def test_paragraph_count_mismatch_leaves_section_untouched(
    tmp_path, stub_planner, new_paragraphs
):
    section = _make_section(tmp_path, ["old one", "old two"])
    stub_planner["plan"] = _plan(new_paragraphs)

    changed, reason = await _apply(section)

    texts = [b.text for b in section.blocks if hasattr(b, "text") and b.text.startswith(("old", "only", "a", "b", "c"))]
    assert changed is False
    assert texts == ["old one", "old two"]
    assert "needs_review" in reason
    assert "model reason" in reason
